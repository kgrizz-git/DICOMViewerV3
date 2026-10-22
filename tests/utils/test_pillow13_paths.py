"""Pillow code paths the app uses, exercised with synthetic in-memory images.

Added for the Pillow 13 bump: PNG/JPEG/TIFF writes, uint8 ``fromarray`` with an
explicit mode, resize/thumbnail, ImageDraw/ImageFont overlay text, ImageOps
invert, PIL-to-QImage conversion, GIF frames, XLSX image embedding, and PDF
output merged by pypdf. No DICOM data is involved.
"""

from __future__ import annotations

import io
from pathlib import Path

import numpy as np
import pytest
from PIL import Image, ImageDraw, ImageFont, ImageOps

from core.lut_display import invert_color_export_image, pil_from_display_array
from utils.image_utils import array_to_image, resize_image


def _gradient_rgb(size: int = 32) -> Image.Image:
    ramp = np.linspace(0, 255, size, dtype=np.uint8)
    arr = np.stack([np.tile(ramp, (size, 1)), np.tile(ramp[:, None], (1, size)), np.full((size, size), 90, np.uint8)], axis=-1)
    return Image.fromarray(arr, mode="RGB")


def test_installed_pillow_is_13_or_newer() -> None:
    assert int(Image.__version__.split(".")[0]) >= 13


@pytest.mark.parametrize("mode", ["L", "RGB"])
def test_png_round_trip_is_lossless(tmp_path: Path, mode: str) -> None:
    image = _gradient_rgb().convert(mode)
    path = tmp_path / "out.png"
    image.save(path, "PNG")
    with Image.open(path) as loaded:
        assert loaded.mode == mode
        assert np.array_equal(np.array(loaded), np.array(image))


@pytest.mark.parametrize("mode", ["L", "RGB"])
def test_jpeg_quality_95_round_trip_is_close(tmp_path: Path, mode: str) -> None:
    image = _gradient_rgb().convert(mode)
    path = tmp_path / "out.jpg"
    image.save(path, "JPEG", quality=95)
    with Image.open(path) as loaded:
        assert loaded.mode == mode
        diff = np.abs(np.array(loaded).astype(int) - np.array(image).astype(int))
        assert diff.mean() < 4


def test_tiff_round_trip_8bit_and_16bit(tmp_path: Path) -> None:
    rgb = _gradient_rgb()
    rgb.save(tmp_path / "rgb.tiff", "TIFF")
    with Image.open(tmp_path / "rgb.tiff") as loaded:
        assert np.array_equal(np.array(loaded), np.array(rgb))

    arr16 = (np.arange(64, dtype=np.uint16).reshape(8, 8) * 1000).astype(np.uint16)
    Image.fromarray(arr16).save(tmp_path / "g16.tiff", "TIFF")
    with Image.open(tmp_path / "g16.tiff") as loaded:
        assert np.array_equal(np.array(loaded), arr16)


def test_16bit_png_needs_explicit_i16_mode(tmp_path: Path) -> None:
    """Pillow 13 no longer clips mode-I PNGs to 16 bits; the app must not save them."""
    arr = (np.arange(16, dtype=np.uint16).reshape(4, 4) * 4000).astype(np.uint16)
    path = tmp_path / "g16.png"
    Image.fromarray(arr).convert("I;16").save(path, "PNG")
    with Image.open(path) as loaded:
        assert np.array_equal(np.array(loaded).astype(np.uint16), arr)


def test_fromarray_explicit_mode_with_uint8_arrays() -> None:
    gray = np.arange(16, dtype=np.uint8).reshape(4, 4)
    rgb = np.zeros((4, 4, 3), dtype=np.uint8)
    assert Image.fromarray(gray, mode="L").mode == "L"
    assert Image.fromarray(rgb, mode="RGB").mode == "RGB"
    assert pil_from_display_array(gray).mode == "L"
    assert pil_from_display_array(rgb).mode == "RGB"
    assert array_to_image(gray) is not None
    assert array_to_image(gray).mode == "L"  # type: ignore[union-attr]
    # The viewer's inversion path: 255 - uint8 stays uint8.
    assert Image.fromarray(255 - gray, mode="L").getpixel((0, 0)) == 255


@pytest.mark.parametrize("mode", ["L", "RGB"])
def test_resize_and_thumbnail_lanczos(mode: str) -> None:
    image = _gradient_rgb(40).convert(mode)
    assert image.resize((80, 80), Image.Resampling.LANCZOS).size == (80, 80)
    thumb = image.copy()
    thumb.thumbnail((16, 16), Image.Resampling.LANCZOS)
    assert thumb.size == (16, 16)
    assert thumb.mode == mode
    resized = resize_image(image, 20, 20)
    assert resized is not None
    assert resized.size == (20, 20)


def test_imagedraw_text_marks_pixels() -> None:
    image = Image.new("RGB", (80, 24), (0, 0, 0))
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default()
    draw.text((2, 2), "R 12", fill=(255, 255, 0), font=font)
    draw.line([(0, 20), (79, 20)], fill=(0, 255, 0), width=1)
    arr = np.array(image)
    assert arr[..., :2].max() == 255
    assert tuple(arr[20, 40]) == (0, 255, 0)


def test_invert_color_export_image_uses_imageops_for_rgb(monkeypatch: pytest.MonkeyPatch) -> None:
    image = Image.new("RGB", (4, 4), (10, 20, 30))
    monkeypatch.setattr("core.dicom_color.is_color_image", lambda _ds: (True, "RGB"))
    out = invert_color_export_image(image, object(), True)
    assert out.getpixel((0, 0)) == (245, 235, 225)
    assert ImageOps.invert(Image.new("L", (2, 2), 5)).getpixel((0, 0)) == 250


def test_pil_array_to_qimage_conversion_matches_pixels() -> None:
    from PySide6.QtGui import QImage

    rgb = np.ascontiguousarray(np.array(_gradient_rgb(16)))
    qimg = QImage(rgb.data, 16, 16, 16 * 3, QImage.Format.Format_RGB888).copy()
    assert qimg.pixelColor(5, 3).getRgb()[:3] == tuple(int(v) for v in rgb[3, 5])

    gray = np.ascontiguousarray(np.array(_gradient_rgb(16).convert("L")))
    qgray = QImage(gray.data, 16, 16, 16, QImage.Format.Format_Grayscale8).copy()
    assert qgray.pixelColor(7, 2).red() == int(gray[2, 7])


def test_multi_frame_gif_save_keeps_frames_and_duration(tmp_path: Path) -> None:
    frames = [Image.new("RGB", (16, 16), (i * 60, 0, 0)) for i in range(4)]
    path = tmp_path / "cine.gif"
    frames[0].save(path, save_all=True, append_images=frames[1:], duration=80, loop=0)
    with Image.open(path) as gif:
        assert gif.n_frames == 4
        assert gif.info["duration"] == 80


def test_openpyxl_embeds_pil_generated_png(tmp_path: Path) -> None:
    from openpyxl import Workbook, load_workbook
    from openpyxl.drawing.image import Image as XLImage

    png = tmp_path / "module.png"
    _gradient_rgb(24).save(png, "PNG")
    wb = Workbook()
    wb.active.add_image(XLImage(str(png)), "B2")
    out = tmp_path / "qa.xlsx"
    wb.save(out)
    assert len(load_workbook(out).active._images) == 1  # pyright: ignore[reportPrivateUsage]


def test_pillow_pdf_pages_merge_with_pypdf(tmp_path: Path) -> None:
    from pypdf import PdfReader, PdfWriter

    paths = []
    for i, colour in enumerate([(200, 30, 30), (30, 30, 200)]):
        path = tmp_path / f"p{i}.pdf"
        Image.new("RGB", (60, 40), colour).save(path, "PDF", resolution=72)
        paths.append(path)
    writer = PdfWriter()
    for path in paths:
        for page in PdfReader(str(path)).pages:
            writer.add_page(page)
    merged = io.BytesIO()
    writer.write(merged)
    merged.seek(0)
    assert len(PdfReader(merged).pages) == 2
